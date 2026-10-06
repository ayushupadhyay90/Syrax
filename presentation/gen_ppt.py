# -*- coding: utf-8 -*-
"""Generate the 10-slide Syrax presentation (theme: #3B82F6 on #040711)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

BG    = RGBColor(0x04, 0x07, 0x11)   # deep space
BLUE  = RGBColor(0x3B, 0x82, 0xF6)   # electric blue (theme)
DEEP  = RGBColor(0x1E, 0x40, 0xAF)   # deep blue
PANEL = RGBColor(0x0A, 0x14, 0x2C)   # panel navy
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xCF, 0xE3, 0xFF)   # pale blue text
MUTED = RGBColor(0x93, 0xC5, 0xFD)   # muted blue
GREY  = RGBColor(0x64, 0x74, 0x8B)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
TOTAL = 10


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    return s


def rect(s, l, t, w, h, fill=BLUE, line=None, rounded=False):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(l), Inches(t), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(1.25)
    shp.shadow.inherit = False
    return shp


def oval(s, l, t, d, fill=None, line=None, lw=2.0):
    shp = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(l), Inches(t), Inches(d), Inches(d))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    return shp


def text(s, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None):
    """runs = [(txt, pt, color, bold), ...] — each its own paragraph."""
    box = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, (txt, pt, color, bold) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        r = p.add_run(); r.text = txt
        r.font.size = Pt(pt); r.font.bold = bold
        r.font.color.rgb = color; r.font.name = 'Calibri'
    return box


def header(s, title, num):
    rect(s, 0.55, 0.52, 0.62, 0.075, BLUE)                 # accent bar
    text(s, 0.5, 0.66, 12.3, 0.95, [(title, 33, WHITE, True)])
    # footer
    text(s, 0.55, 7.02, 4.0, 0.35, [('SYRAX — Mark 1', 10, GREY, False)])
    text(s, 11.6, 7.02, 1.3, 0.35, [(f'{num} / {TOTAL}', 10, GREY, False)], align=PP_ALIGN.RIGHT)
    # decorative glow circles top-right
    oval(s, 11.9, -1.1, 2.6, fill=None, line=DEEP, lw=6)
    oval(s, 12.55, 0.35, 0.55, fill=BLUE)


def bullets(s, items, top=1.9, left=0.6, width=12.1, size=17, gap=13):
    box = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(4.6))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if isinstance(item, tuple):     # (text, color)
            txt, color = item
        else:
            txt, color = item, LIGHT
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        r = p.add_run(); r.text = '•  ' + txt
        r.font.size = Pt(size); r.font.color.rgb = color
        r.font.name = 'Calibri'
    return box


def pipeline(s, labels, top=5.45, boxw=2.15, boxh=0.78):
    """Row of rounded boxes with blue arrows between them."""
    n = len(labels)
    gap = 0.30
    total_w = n * boxw + (n - 1) * (gap + 0.32)
    left = (13.333 - total_w) / 2
    for i, lab in enumerate(labels):
        x = left + i * (boxw + gap + 0.32)
        b = rect(s, x, top, boxw, boxh, PANEL, BLUE, rounded=True)
        tf = b.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = lab
        r.font.size = Pt(12); r.font.bold = True
        r.font.color.rgb = LIGHT; r.font.name = 'Calibri'
        if i < n - 1:
            text(s, x + boxw + 0.01, top + 0.13, 0.34, 0.5,
                 [('›', 22, BLUE, True)], align=PP_ALIGN.CENTER)


def notes(s, txt):
    s.notes_slide.notes_text_frame.text = txt


# ─────────────────────────── SLIDE 1 · TITLE ───────────────────────────
s = slide()
oval(s, 9.6, 0.7, 5.4, fill=None, line=DEEP, lw=8)
oval(s, 11.9, 4.9, 2.2, fill=None, line=PANEL, lw=10)
rect(s, 0.9, 1.55, 1.9, 0.5, None, BLUE, rounded=True)
text(s, 0.9, 1.6, 1.9, 0.4, [('MARK 1', 16, BLUE, True)], align=PP_ALIGN.CENTER)
text(s, 0.82, 2.1, 9.5, 1.7, [('SYRAX', 110, WHITE, True)])
rect(s, 0.9, 3.9, 3.4, 0.07, BLUE)
text(s, 0.9, 4.12, 11.0, 1.5, [
    ('Voice-Driven 3D Particle Web UI Agent', 30, LIGHT, True),
    ('Speak  ·  Search  ·  Play — in real browser tabs', 18, MUTED, False),
], spacing=8)
text(s, 0.9, 6.35, 11.5, 0.9, [
    ('Created by Ayush, Navam & Sandeep', 15, WHITE, True),
    ('Live:  https://ayushupadhyay90.github.io/Syrax/', 14, BLUE, False),
], spacing=4)
notes(s, 'Syrax — a voice agent you talk to like a person. It listens once, acts in real '
         'Chrome tabs, and installs as an app on desktop and Android. Creators: Ayush, Navam & Sandeep.')

# ─────────────────────── SLIDE 2 · WHAT IS SYRAX ───────────────────────
s = slide(); header(s, 'What is Syrax?', 2)
bullets(s, [
    'Voice-controlled AI assistant that runs entirely in your browser',
    '3D particle interface reacts to your voice and to every action',
    'Operates REAL Chrome tabs — never an embedded fake panel',
    'Installable on desktop + Android from a single link (PWA)',
    'Brain: DeepSeek deepseek-flash  ·  Voice: Web Speech STT + fast TTS',
    'Answers questions, opens pages, searches the web, plays music — hands-free',
], top=1.95, size=18.5, gap=16)
notes(s, 'Positioning: a personal browser agent. Everything it does happens in real tabs, '
         'so you always see and control the result. One responsive web app for laptop and phone.')

# ─────────────────────── SLIDE 3 · CORE FEATURES ───────────────────────
s = slide(); header(s, 'Core Features', 3)
bullets(s, [
    'One-shot mic: listens for one command, then ALWAYS turns itself off',
    'Instant halt: “stop Syrax”, “stop”, “shut up” — everything stops at once',
    'Questions are heard but never executed — clear commands only',
    '“Play a song” → asks which song → your answer plays the direct video',
    'Search chains wait for your FULL sentence before acting (no early fires)',
    'Zero animation noise while executing — quiet status text, then the tab opens',
    'Mic never hears video/audio — all media is muted while listening',
], top=1.9, size=17.5, gap=13)
notes(s, 'These all came directly from user feedback: safety (no accidental actions), '
         'speed (no wasted animation), and completeness (never cut a sentence short).')

# ──────────────────────── SLIDE 4 · ARCHITECTURE ───────────────────────
s = slide(); header(s, 'System Architecture', 4)
bullets(s, [
    'Frontend: React + Vite + Tailwind CSS  —  fast, minimal, themeable',
    '3D layer: Three.js particle system (vendor-split for a light main bundle)',
    'Brain: DeepSeek API — JSON mode, thinking off, temperature 0.5',
    'Voice: Web Speech API for STT; layered TTS (fish → pocket → browser voice)',
    'Browser ops: popup-proof reserved tab + real window navigation',
    'Storage: localStorage — chat history + one-time API key per device',
], top=1.85, size=17, gap=12)
pipeline(s, ['Mic / STT', 'Intent Gate', 'DeepSeek', 'Tab Ops', 'TTS'], top=5.5)
notes(s, 'Two hard rules shape the architecture: local intent handles anything that acts '
         '(instant and safe), the LLM only ever chats and answers — it cannot open or play anything.')

# ──────────────────────── SLIDE 5 · VOICE PIPELINE ─────────────────────
s = slide(); header(s, 'The Voice Pipeline', 5)
bullets(s, [
    ('1.  Mic tap reserves a REAL tab synchronously — inside the user gesture', WHITE),
    ('2.  STT buffers your speech; fires only once the sentence is complete', LIGHT),
    ('3.  Local intent parser matches the command instantly (clear commands)', LIGHT),
    ('4.  LLM answers questions — but is never allowed to execute actions', LIGHT),
    ('5.  Action navigates the reserved tab — popup blockers can’t stop it', LIGHT),
    ('6.  TTS replies at 1.1× speed, then the mic returns to OFF (one-shot)', LIGHT),
], top=1.9, size=17.5, gap=15)
pipeline(s, ['Mic tap', 'Listen', 'Decide', 'Act in tab', 'Reply'], top=5.55)
notes(s, 'The reserve step is the popup trick: Chrome only allows window.open for ~5 seconds '
         'after a real click. We open the tab at mic-click and only navigate it later, so voice '
         'commands can never be blocked. Dangling-word flush prevents mid-sentence firing.')

# ───────────────────── SLIDE 6 · REAL BROWSER OPS ─────────────────────
s = slide(); header(s, 'Real Browser Operations', 6)
bullets(s, [
    'Every operation opens in a REAL Chrome tab — user requirement, always visible',
    'Popup-blocker-proof: tab reserved at mic-click, commands only navigate it',
    'Play → direct YouTube watch URL via a 9-source parallel resolver (~2 s)',
    'Resolver slow or down? → YouTube results page fallback — the tab ALWAYS opens',
    'Searches continue in the already-open tab — zero tab stacking',
    '“Open YouTube” always lands on the homepage — never stale content',
    'Now-playing strip appears while music plays, clears when you close the tab',
], top=1.9, size=17, gap=12)
notes(s, 'The resolver races nine public sources (invidious, proxies, jina) client-side because '
         'serverless YouTube APIs are gone on GitHub Pages. First valid video ID wins; failure '
         'still lands you on YouTube search — the operation never hangs.')

# ────────────────────── SLIDE 7 · DESIGN & EXPERIENCE ──────────────────
s = slide(); header(s, 'Design & Experience', 7)
bullets(s, [
    'Theme: electric blue #3B82F6 on deep-space #040711 — one coherent palette',
    'Bold SYRAX wordmark + “SYRAX — Mark 1” badge (STT mishear “cyrex” is auto-fixed)',
    'Majestic chat-bubble crown logo — generated as scalable, maskable PWA icon',
    'Glow language: live status line, mic rings, pulsing side panels',
    'Android-first responsive layout: stacking grid, capped lists, sized hero',
    'Minimal UI: two slim panels, one conversation, nothing that competes with voice',
], top=1.9, size=17.5, gap=14)
# theme swatches
for i, (c, lab) in enumerate([(BG, '#040711'), (DEEP, '#1E40AF'), (BLUE, '#3B82F6'), (LIGHT, '#CFE3FF')]):
    x = 0.65 + i * 2.6
    rect(s, x, 5.9, 2.2, 0.62, c, GREY if c == BG else None, rounded=True)
    text(s, x, 6.58, 2.2, 0.35, [(lab, 11.5, MUTED, False)], align=PP_ALIGN.CENTER)
notes(s, 'The palette was swept across every screen and computed styles verified: zero leftover '
         'colors, 45 themed elements. The logo is drawn by a custom script — no image assets needed.')

# ─────────────────────── SLIDE 8 · PRIVACY & SECURITY ──────────────────
s = slide(); header(s, 'Privacy & Security — Honest', 8)
bullets(s, [
    'No false privacy claims anywhere — all “privacy-first” wording removed',
    'The truth, stated in-app: DeepSeek API processes your prompts (provider storage)',
    'API key is entered once per device and lives only in that browser’s localStorage',
    'Repo stays public and key-free; ?key= links save the key, then strip themselves',
    'One-shot mic — Syrax never keeps listening in the background',
    'All media muted while listening: the mic hears you, never the video',
], top=1.9, size=17.5, gap=14)
notes(s, 'Trust beats marketing: we say exactly what leaves the browser (prompt text to DeepSeek) '
         'instead of claiming privacy we can’t guarantee. No secrets are ever committed to the repo.')

# ─────────────────────── SLIDE 9 · DEPLOY & VERIFY ─────────────────────
s = slide(); header(s, 'Deploy & Verification', 9)
bullets(s, [
    'Push to main → GitHub Actions builds → GitHub Pages auto-deploys',
    'PWA: manifest + service worker, relative paths — installs on desktop & Android',
    'End-to-end verified LIVE: play → direct video · open → homepage · search → Google',
    'Also verified: two-step song ask, halt commands, now-playing strip lifecycle',
    'Hashed bundles — if a change looks stale, one Ctrl+F5 loads the latest',
    'Stack: Vite · React · Three.js · Tailwind · DeepSeek · GitHub Pages',
], top=1.9, size=17.5, gap=14)
notes(s, 'Everything on this deck was tested against the deployed site, not just locally — '
         'commands typed and results read back from real tabs before shipping each change.')

# ─────────────────────── SLIDE 10 · LINKS & ROADMAP ────────────────────
s = slide(); header(s, 'Links & Roadmap', 10)
# left panel — links
rect(s, 0.6, 1.95, 6.0, 3.0, PANEL, BLUE, rounded=True)
text(s, 0.95, 2.2, 5.4, 2.6, [
    ('USE IT NOW', 14, BLUE, True),
    ('Desktop app', 15, WHITE, True),
    ('https://ayushupadhyay90.github.io/Syrax/', 13.5, LIGHT, False),
    ('Mobile (Android)', 15, WHITE, True),
    ('Same link → Chrome ⋮ → “Add to Home screen”', 13.5, LIGHT, False),
], spacing=6)
# right panel — roadmap
rect(s, 6.9, 1.95, 5.85, 3.0, PANEL, DEEP, rounded=True)
text(s, 7.25, 2.2, 5.2, 2.6, [
    ('ROADMAP', 14, MUTED, True),
    ('•  Fish Audio voice clone (JARVIS voice)', 14.5, LIGHT, False),
    ('•  Faster STT tuning + more local actions', 14.5, LIGHT, False),
    ('•  Longer memory & multi-step tasks', 14.5, LIGHT, False),
    ('•  Widget / always-on assistant mode', 14.5, LIGHT, False),
], spacing=8)
text(s, 0.6, 5.4, 12.1, 1.3, [
    ('Thank you', 40, WHITE, True),
    ('Syrax — Mark 1  ·  Created by Ayush, Navam & Sandeep', 16, MUTED, False),
], spacing=6)
notes(s, 'One URL for both experiences: the same responsive PWA runs on desktop and installs '
         'on Android. Questions?')

prs.save(r'C:\Users\Ayush\OneDrive\Desktop\Syrax\presentation\Syrax-Presentation.pptx')
print('saved', len(prs.slides.__iter__.__self__._sldIdLst), 'slides')
